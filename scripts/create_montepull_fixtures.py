"""Crée des copies réduites et anonymisées des Excel Montepull pour les tests."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

OUT = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "montepull"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Commandes"
    ws.append(
        [
            "Commande",
            "Client",
            "Article",
            "Designation",
            "Qte_Commandee",
            "Qte_Lancee",
            "Somme_Operations",
            "Somme_Op_70_Commande",
            "Somme_Op_70_Article",
            "Derniere_Operation",
            "Qte_Fabriquee_Derniere_Operation",
        ]
    )
    ws.append(["SO_TEST_001", "CL_TEST", "ART_T01", "Article test XS", 100, 0, 0, 0, 0, None, 0])
    ws.append(["SO_TEST_001", "CL_TEST", "ART_T02", "Article test S", 200, 50, 10, 0, 0, 30, 40])
    ws.append(["SO_TEST_001", "CL_TEST", "ART_T03", "Article test M", 80, 80, 80, 80, 80, 70, 80])
    wb.save(OUT / "commandes_sample.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Details"
    ws.append(
        [
            "N° OF",
            "N° CMD",
            "Fournisseur",
            "Article",
            "Désignation",
            "Opération",
            "Date Début",
            "Date Fin",
            "Quantité Réelle",
            "Quantité Rejet",
            "Qte OF",
            "Poids",
            "CMD Date",
            "CMD Client",
            "CMD Client Name",
            "CMD Date Livraison",
            "N CMD Client",
            "Atelier",
            "Statut",
            "Npaquet",
            "Date Opération",
        ]
    )
    ws.append(
        [
            "OF_TEST_001",
            "SO_TEST_001",
            "-",
            "ART_T02",
            "Article test S",
            30,
            datetime(2026, 7, 1, 10, 0, 0),
            datetime(2026, 7, 1, 10, 0, 2),
            20,
            0,
            50,
            0,
            date(2026, 4, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 6, 30),
            "REF1",
            "Montepull",
            "Terminé",
            "1",
            date(2026, 7, 1),
        ]
    )
    ws.append(
        [
            "OF_TEST_001",
            "SO_TEST_001",
            "-",
            "ART_T02",
            "Article test S",
            30,
            datetime(2026, 7, 1, 11, 0, 0),
            datetime(2026, 7, 1, 11, 0, 1),
            15,
            1,
            50,
            0,
            date(2026, 4, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 6, 30),
            "REF1",
            "Montepull",
            "Terminé",
            "2",
            date(2026, 7, 1),
        ]
    )
    ws.append(
        [
            "OF_TEST_001",
            "SO_TEST_001",
            "-",
            "ART_T02",
            "Article test S",
            40,
            datetime(2026, 7, 2, 9, 0, 0),
            datetime(2026, 7, 2, 10, 30, 0),
            35,
            0,
            50,
            0,
            date(2026, 4, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 6, 30),
            "REF1",
            "Montepull",
            "En cours",
            "1",
            date(2026, 7, 2),
        ]
    )
    ws.append(
        [
            "OF_TEST_002",
            "SO_TEST_002",
            "-",
            "ART_T03",
            "Article test M",
            70,
            datetime(2026, 7, 3, 8, 0, 0),
            datetime(2026, 7, 3, 8, 0, 5),
            80,
            0,
            80,
            0,
            date(2026, 5, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 7, 1),
            "REF2",
            "Montepull",
            "Terminé",
            "1",
            date(2026, 7, 3),
        ]
    )
    ws.append(
        [
            "OF_TEST_001",
            "SO_TEST_001",
            "-",
            "ART_T02",
            "Article test S",
            35,
            datetime(2026, 7, 1, 12, 0, 0),
            datetime(2026, 7, 1, 12, 0, 1),
            5,
            0,
            50,
            0,
            date(2026, 4, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 6, 30),
            "REF1",
            "Montepull",
            "Terminé",
            "3",
            date(2026, 7, 1),
        ]
    )
    ws.append(
        [
            "OF_TEST_003",
            "SO_TEST_003",
            "-",
            "ART_T01",
            "Article test XS",
            10,
            datetime(2026, 7, 4, 12, 0, 0),
            datetime(2026, 7, 4, 11, 0, 0),
            10,
            0,
            100,
            0,
            date(2026, 5, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 8, 1),
            "REF3",
            "Montepull",
            "Terminé",
            "1",
            date(2026, 7, 4),
        ]
    )

    ws2 = wb.create_sheet("Par Operation")
    ws2.append(
        [
            "N° OF",
            "N° CMD",
            "Fournisseur",
            "Article",
            "Désignation",
            "Opération",
            "Qte OF",
            "CMD Date",
            "CMD Client",
            "CMD Client Name",
            "CMD Date Livraison",
            "N CMD Client",
            "Atelier",
            "Quantité Réelle",
            "Quantité Rejet",
        ]
    )
    ws2.append(
        [
            "OF_TEST_001",
            "SO_TEST_001",
            "-",
            "ART_T02",
            "Article test S",
            30,
            50,
            date(2026, 4, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 6, 30),
            "REF1",
            "Montepull",
            35,
            1,
        ]
    )

    ws3 = wb.create_sheet("Feuille Route")
    ws3.append(
        [
            "N° OF",
            "N° CMD",
            "Fournisseur",
            "Article",
            "Désignation",
            "Qte OF",
            "CMD Date",
            "CMD Client",
            "CMD Client Name",
            "CMD Date Livraison",
            "N CMD Client",
            "Atelier",
            "Op10",
            "Op20",
            "Op30",
            "Op40",
            "Op50",
            "Op60",
            "Op70",
        ]
    )
    ws3.append(
        [
            "OF_TEST_001",
            "SO_TEST_001",
            "-",
            "ART_T02",
            "Article test S",
            50,
            date(2026, 4, 1),
            "CL_TEST",
            "CLIENT TEST",
            date(2026, 6, 30),
            "REF1",
            "Montepull",
            0,
            0,
            35,
            35,
            0,
            0,
            0,
        ]
    )
    wb.save(OUT / "ListeSuivi_sample.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Details"
    ws.append(
        [
            "Opération",
            "N° OF",
            "N° CMD",
            "Article",
            "Désignation",
            "Qté commandée",
            "Quantité Réelle",
            "Quantité Rejet",
            "Poids",
            "Statut",
            "Date Début",
            "Date Fin",
        ]
    )
    ws.append([30, "OF_TEST_001", "SO_TEST_001", "ART_T02", "Article test S", 200, 20, 0, 0, "Terminé", "2026-07-01 10:00:00", "2026-07-01 10:00:02"])
    ws.append([30, "OF_TEST_001", "SO_TEST_001", "ART_T02", "Article test S", 200, 20, 0, 0, "Terminé", "2026-07-01 11:00:00", "2026-07-01 12:00:00"])
    ws.append([40, "OF_TEST_001", "SO_TEST_001", "ART_T02", "Article test S", 200, 35, 0, 0, "Terminé", "2026-07-02 09:00:00", "2026-07-02 09:00:01"])
    ws.append([10, "OF_TEST_003", "SO_TEST_003", "ART_T01", "Article test XS", 100, 10, 0, 0, "Terminé", "2026-07-04 12:00:00", "2026-07-04 11:00:00"])
    ws.append([70, "OF_TEST_002", "SO_TEST_002", "ART_T03", "Article test M", 80, 80, -1, 0, "Terminé", "2026-07-03 08:00:00", "2026-07-03 08:00:05"])
    ws2 = wb.create_sheet("Par OF")
    ws2.append(
        [
            "Opération",
            "N° OF",
            "N° CMD",
            "Article",
            "Désignation",
            "Qté commandée",
            "Quantité Réelle",
            "Quantité Rejet",
            "Poids",
            "Statut",
            "Date Début",
            "Date Fin",
        ]
    )
    ws2.append([30, "OF_TEST_001", "SO_TEST_001", "ART_T02", "Article test S", 200, 40, 0, 0, "Terminé", "2026-07-01 10:00:00", "2026-07-01 12:00:00"])
    wb.save(OUT / "SuiviOperations_sample.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Marwa simple Standard"
    ws["C4"] = "Calcul du besoin d'achat de la reference :"
    ws["G4"] = "DOUBLYGILF_TEST"
    ws["B6"] = "Delai client"
    ws["C6"] = "Client"
    ws["D6"] = "Ref"
    ws["E6"] = "Jauge"
    ws["F6"] = "Version"
    ws["B7"] = date(2026, 10, 9)
    ws["C7"] = "CLIENT_TEST"
    ws["D7"] = "DOUBLYGILF_TEST"
    ws["E7"] = 7
    ws["F7"] = "CREME"
    ws["F9"] = "Taille"
    for col, size in zip("GHIJK", ["XS", "S", "M", "L", "XL"]):
        ws[f"{col}9"] = size
    ws["L9"] = "Total Besoin"
    ws["F12"] = "Total comde"
    for col, qty in zip("GHIJK", [10, 20, 30, 15, 10]):
        ws[f"{col}12"] = qty
    ws["L12"] = 85
    ws["F13"] = "Poids"
    for col in "GHIJK":
        ws[f"{col}13"] = 0.3
    ws["B14"] = "Fournisseur"
    ws["C14"] = "Composition"
    ws["F14"] = "Besoin fil"
    ws["B15"] = "BVB"
    ws["C15"] = "63%ACR"
    ws["F15"] = 1
    for col, qty in zip("GHIJK", [3, 6, 9, 4.5, 3]):
        ws[f"{col}15"] = qty
    ws["L15"] = 25.5
    ws["B21"] = "SML"
    ws["C21"] = "Vignette de composition"
    ws["F21"] = 0.26
    ws["G21"] = 100
    ws["B22"] = "SML"
    ws["C22"] = "RFID"
    ws["F22"] = 0.8
    ws["G22"] = 100
    wb.create_sheet("Fiche de prix")
    wb.save(OUT / "DOUBLYGILF_sample.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Feuil3"
    ws["A2"] = "FICHE NOMENCLATURE"
    ws["A4"] = "Client :"
    ws["B4"] = "CLIENT_TEST"
    ws["A5"] = "Reference :"
    ws["B5"] = "DOULBYGILF_TEST"
    ws["A8"] = "Designation : PULL"
    ws["A15"] = "FOURNITURES"
    ws["D15"] = "Reference"
    ws["F15"] = "Emploi"
    ws["G15"] = "Besoin"
    ws["H15"] = "Fournisseur"
    ws["I15"] = "Prix"
    ws["A16"] = "Vignette de composition"
    ws["F16"] = 2
    ws["H16"] = "SML"
    ws["I16"] = 0.2676
    ws["A17"] = "RFID"
    ws["F17"] = 1
    ws["H17"] = "SML"
    ws["I17"] = 0.8028
    ws["A18"] = "Sachet"
    ws["F18"] = 1
    ws["H18"] = "MARPLAST"
    ws["I18"] = 0.56
    ws["J18"] = "#VALUE!"
    wb.save(OUT / "Nomenclature_DOULBYGILF_sample.xlsx")

    print(f"Fixtures written to {OUT}")
    for p in sorted(OUT.glob("*.xlsx")):
        print(f"  {p.name} ({p.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
